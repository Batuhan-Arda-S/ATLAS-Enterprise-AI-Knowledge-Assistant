import pytest
import asyncio
from app.services.file_storage import file_storage
from app.services.file_upload import file_upload_service
from app.services.document_parser import document_parser
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
from docx import Document as DocxDocument
from pypdf import PdfWriter
from sqlalchemy import select
from app.models.document_chunk import DocumentChunk

@pytest.fixture
def temp_storage(monkeypatch,tmp_path):
    monkeypatch.setattr(file_storage,"storage_root",tmp_path)
    return tmp_path


def valid_pdf_bytes():
    stream = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.write(stream)
    return stream.getvalue()


def test_upload_txt_file(client,temp_storage):
    text = "Çalışanların yıllık izin politikası."
    file_content=text.encode("utf-8")
    response = client.post("/documents/upload",
                           data={"title":"AI Notes","department":"Engineering"},
                           files={"file":("notes.txt",file_content,"text/plain")})

    assert response.status_code == 201

    response_data=response.json()

    assert response_data["title"] == "AI Notes"
    assert response_data["department"] == "Engineering"
    assert response_data["filename"] == "notes.txt"
    assert response_data["content_type"] == "text/plain"
    assert response_data["file_size"] == len(file_content)
    assert response_data["processing_status"] == "processed"
    assert isinstance(response_data["id"], int)

    documents_directory=(temp_storage/"documents")
    stored_files=list(documents_directory.iterdir())
    assert len(stored_files) == 1
    assert(stored_files[0].read_bytes()==file_content)

    document_response=client.get(f"/documents/{response_data['id']}")

    assert(document_response.json()["content"] == text)

    assert document_response.status_code == 200
    assert document_response.json()["title"] == "AI Notes"

def test_upload_unsupported_file_type(client):
    response = client.post("/documents/upload",
                           data={"title": "Image", "department": "Engineering"},
                           files={"file":("image.jpg",b"fake image content","image/jpeg")})

    assert response.status_code == 415

    assert response.json() == {
        "detail": "Unsupported file type"
    }

def test_upload_empty_file(client):
    response = client.post("/documents/upload",
                           data={"title": "AI Notes", "department": "Engineering"},
                           files={"file":("empty.txt",b"","text/plain")})

    assert response.status_code == 400

    assert response.json() == {
        "detail": "File cannot be empty"
    }

def test_upload_too_large_file(client):
    large_content=b"a" *(10*1024*1024+1)

    response = client.post("/documents/upload",
                           data={"title": "AI Notes", "department": "Engineering"},
                           files={"file":("large.txt",large_content,"text/plain")})

    assert response.status_code == 413

    assert response.json() == {
        "detail": "File is too large"
    }

def test_upload_without_file(client):
    response = client.post("/documents/upload",
                           data={"title":"Missing File","department":"Engineering"},)

    assert response.status_code == 422

def test_pdf_upload_is_processed(client,temp_storage,monkeypatch):

    extracted_text=("This text came from the PDF parser.")

    monkeypatch.setattr(document_parser,"extract_pdf_text",lambda contents:extracted_text)

    response = client.post("/documents/upload",data={"title":"Remote Work Policy",
                                                     "department":"Human Resources",},
                                               files={"file":
                                                          ("policy.pdf",
                                                           valid_pdf_bytes(),
                                                           "application/pdf")},)

    assert response.status_code == 201

    response_data=response.json()

    assert (response_data["processing_status"]=="processed")

    document_response=client.get(f"/documents/{response_data['id']}")

    assert document_response.status_code == 200

    assert (document_response.json()["content"] == extracted_text)


def test_docx_upload_is_processed(client,temp_storage):

    docx_document=DocxDocument()

    docx_document.add_paragraph("Enterprise AI Knowledge Assistant")
    docx_document.add_paragraph("Employees can work remotely two days per week.")

    docx_stream=BytesIO()
    docx_document.save(docx_stream)
    docx_stream.seek(0)

    docx_content=docx_stream.read()

    response = client.post("/documents/upload",data={"title":"Remote Work Policy",
                                                     "department":"Human Resources",},
                                               files={"file":("remote_work_policy.docx",
                                                              docx_content,
                                                              ("application/vnd.openxmlformats-officedocument."
                                                               "wordprocessingml.document"))})
    assert response.status_code == 201

    response_data=response.json()

    assert (response_data["processing_status"]=="processed")

    document_response=client.get(f"/documents/{response_data['id']}")

    assert document_response.status_code == 200

    content=document_response.json()["content"]

    assert("Enterprise AI Knowledge Assistant" in content)

    assert ("Employees can work remotely two days per week." in content)


def test_pdf_parsing_failure_marks_document_failed(client,temp_storage,monkeypatch):

    original_pdf = valid_pdf_bytes()
    def raise_parsing_error(contents:bytes):
        raise ValueError("PDF parsing failed")

    monkeypatch.setattr(document_parser,"extract_pdf_text",raise_parsing_error)
    response = client.post("/documents/upload",data={"title":"Broken Policy",
                                                     "department":"Human Resources",},
                                               files={"file":("broken.pdf",
                                                          original_pdf,
                                                              "application/pdf")},)

    response_data=response.json()
    assert (response_data["processing_status"]=="failed")

    document_response=client.get(f"/documents/{response_data['id']}")
    assert document_response.status_code == 200

    assert (document_response.json()["content"]=="")

    documents_directory=(temp_storage/"documents")

    stored_files=list(documents_directory.iterdir())
    assert len(stored_files)==1

    assert stored_files[0].read_bytes() == original_pdf

def test_txt_upload_creates_document_chunks(client,temp_storage,db_session):

    sentence=("Çalışanlar uzaktan çalışma için "  "yönetici onayı almalıdır. ")

    text=sentence*20

    response=client.post("/documents/upload",data={"title":"Remote Work Rules",
                                                   "department":"Human Resources",},
                                             files={"file":("rules.txt",
                                                            text.encode("utf-8"),
                                                            "text/plain")})
    assert response.status_code == 201
    response_data=response.json()
    assert response_data["processing_status"]=="processed"
    document_id=response_data["id"]

    statement=(select(DocumentChunk).where(DocumentChunk.document_id == document_id).order_by(DocumentChunk.chunk_index))

    stored_chunks=db_session.scalars(statement).all()

    assert len(stored_chunks)>1

    assert[chunk.chunk_index for chunk in stored_chunks]==list(range(len(stored_chunks)))

    assert all(chunk.document_id == document_id for chunk in stored_chunks)
    assert all(chunk.embedding is not None for chunk in stored_chunks)

    assert all(0<len(chunk.content)<=500 for chunk in stored_chunks)

    document_response=client.get(f"/documents/{document_id}")

    assert document_response.status_code == 200
    assert document_response.json()["content"]==text

def test_whitespace_txt_upload_is_marked_failed(client,temp_storage,db_session):
    response=client.post("/documents/upload",data={"title":"Empty Notes","department":"Engineering"},
                                             files={"file":("notes.txt",b"  \n  ","text/plain")})

    assert response.status_code == 201
    response_data=response.json()
    assert response_data["processing_status"]=="failed"

    document_id=response_data["id"]

    statement=select(DocumentChunk).where(DocumentChunk.document_id == document_id)

    stored_chunks=db_session.scalars(statement).all()

    assert stored_chunks==[]

    document_response=client.get(f"/documents/{document_id}")

    assert document_response.status_code == 200
    assert document_response.json()["content"]==""


def test_embedding_failure_marks_upload_failed(client,temp_storage,db_session,monkeypatch):
    class BrokenEmbedding:
        def embed_passages(self,texts):
            raise RuntimeError("Embedding unavailable")

    monkeypatch.setattr("app.services.document.get_embedding_service",lambda: BrokenEmbedding())
    response = client.post("/documents/upload",data={"title":"Broken embeddings",
                          "department":"Engineering"},
                           files={"file":("notes.txt",b"Some useful policy text","text/plain")})
    assert response.status_code == 201
    assert response.json()["processing_status"] == "failed"
    document_id = response.json()["id"]
    assert db_session.scalars(select(DocumentChunk).where(
        DocumentChunk.document_id == document_id)).all() == []


def test_delete_uploaded_document_removes_its_stored_file(client,temp_storage):
    response = client.post("/documents/upload",data={"title":"Temporary file",
                          "department":"Engineering"},
                           files={"file":("temporary.txt",b"Temporary policy text","text/plain")})
    document_id = response.json()["id"]
    stored_files = list((temp_storage/"documents").iterdir())
    assert len(stored_files) == 1
    assert client.delete(f"/documents/{document_id}").status_code == 200
    assert not stored_files[0].exists()


def test_upload_reads_only_up_to_size_limit():
    class ControlledUpload:
        content_type = "text/plain"
        filename = "large.txt"

        async def read(self, size=-1):
            assert size > 0
            return b"x" * size

    with pytest.raises(Exception) as error:
        asyncio.run(file_upload_service.validate_and_read(ControlledUpload()))
    assert type(error.value).__name__ == "FileTooLargeError"


@pytest.mark.parametrize("filename,content_type,contents", [
    ("fake.pdf", "application/pdf", b"<html>not a pdf</html>"),
    ("fake.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", b"not a zip"),
    ("binary.txt", "text/plain", b"\x00\x01\xff"),
])
def test_upload_rejects_content_type_spoofing(client,filename,content_type,contents):
    response = client.post("/documents/upload", data={"title": "Bad", "department": "HR"},
                           files={"file": (filename, contents, content_type)})
    assert response.status_code == 415


def test_upload_rejects_large_decompressed_docx(client):
    stream = BytesIO()
    with ZipFile(stream, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "x" * (21 * 1024 * 1024))
    response = client.post("/documents/upload", data={"title": "Bomb", "department": "HR"},
                           files={"file": ("bomb.docx", stream.getvalue(),
                                           "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert response.status_code == 413


def test_upload_strips_path_components_from_filename(client,temp_storage):
    response = client.post("/documents/upload", data={"title": "Safe", "department": "HR"},
                           files={"file": ("..\\..\\private.txt", b"Policy text", "text/plain")})
    assert response.status_code == 201
    assert response.json()["filename"] == "private.txt"
    assert len(list((temp_storage / "documents").iterdir())) == 1



